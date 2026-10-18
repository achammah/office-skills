#!/usr/bin/env python3
"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file and reports formula errors.

Default path: paper-xlsx's oracle (openpyxl.oracle). LibreOffice recalculates a
temporary copy; the computed values are spliced into the ORIGINAL package, so
charts, images, OLE objects, comments and every other part are kept.
Fallback (only when openpyxl.oracle is missing, i.e. stock openpyxl): the old
LibreOffice macro path, which re-saves the whole file through LibreOffice and
can drop parts. The JSON output says which path ran.
"""

import json
import sys
import subprocess
import os
import platform
import tempfile
from pathlib import Path
from openpyxl import load_workbook

EXCEL_ERRORS = ['#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A']


def summarize(locations_by_type, total_formulas, method):
    total_errors = sum(len(v) for v in locations_by_type.values())
    result = {
        'status': 'success' if total_errors == 0 else 'errors_found',
        'total_errors': total_errors,
        'error_summary': {},
        'total_formulas': total_formulas,
        'method': method,
    }
    for err_type, locations in locations_by_type.items():
        if locations:
            result['error_summary'][err_type] = {
                'count': len(locations),
                'locations': locations[:20]  # Show up to 20 locations
            }
    return result


def recalc_oracle(filename, timeout, oracle):
    """Recalculate via paper-xlsx's oracle and replace the file atomically."""
    from openpyxl.errors import PaperRefusal

    path = Path(filename).absolute()
    fd, tmp = tempfile.mkstemp(prefix='.recalc-', suffix=path.suffix, dir=path.parent)
    os.close(fd)
    os.unlink(tmp)
    try:
        r = oracle.recalc(str(path), output_path=tmp, timeout=timeout)
        os.replace(tmp, path)
    except PaperRefusal as e:
        return {'error': f'{type(e).__name__}: {e}', 'refusal_kind': getattr(e, 'kind', None),
                'file_unchanged': True, 'method': 'paper-oracle'}
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)

    by_type = {err: [] for err in EXCEL_ERRORS}
    for e in r.errors:
        by_type.setdefault(e['value'], []).append(f"{e['sheet']}!{e['cell']}")
    result = summarize(by_type, r.formula_cells, 'paper-oracle')
    result['cells_written'] = r.cells_written
    if r.pivot_refreshes:
        result['pivot_refreshes'] = r.pivot_refreshes
    return result


def setup_libreoffice_macro():
    """Setup LibreOffice macro for recalculation if not already configured"""
    if platform.system() == 'Darwin':
        macro_dir = os.path.expanduser('~/Library/Application Support/LibreOffice/4/user/basic/Standard')
    else:
        macro_dir = os.path.expanduser('~/.config/libreoffice/4/user/basic/Standard')

    macro_file = os.path.join(macro_dir, 'Module1.xba')

    if os.path.exists(macro_file):
        with open(macro_file, 'r') as f:
            if 'RecalculateAndSave' in f.read():
                return True

    if not os.path.exists(macro_dir):
        subprocess.run(['soffice', '--headless', '--terminate_after_init'],
                      capture_output=True, timeout=10)
        os.makedirs(macro_dir, exist_ok=True)

    macro_content = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>'''

    try:
        with open(macro_file, 'w') as f:
            f.write(macro_content)
        return True
    except Exception:
        return False


def recalc_legacy(filename, timeout):
    """Old path: LibreOffice re-saves the whole file (may drop parts)."""
    abs_path = str(Path(filename).absolute())

    if not setup_libreoffice_macro():
        return {'error': 'Failed to setup LibreOffice macro'}

    cmd = [
        'soffice', '--headless', '--norestore',
        'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application',
        abs_path
    ]

    # Handle timeout command differences between Linux and macOS
    if platform.system() != 'Windows':
        timeout_cmd = 'timeout' if platform.system() == 'Linux' else None
        if platform.system() == 'Darwin':
            try:
                subprocess.run(['gtimeout', '--version'], capture_output=True, timeout=1, check=False)
                timeout_cmd = 'gtimeout'
            except (FileNotFoundError, subprocess.TimeoutExpired):
                pass
        if timeout_cmd:
            cmd = [timeout_cmd, str(timeout)] + cmd

    result = subprocess.run(cmd, capture_output=True, text=True)

    if result.returncode != 0 and result.returncode != 124:  # 124 is timeout exit code
        error_msg = result.stderr or 'Unknown error during recalculation'
        if 'Module1' in error_msg or 'RecalculateAndSave' not in error_msg:
            return {'error': 'LibreOffice macro not configured properly'}
        return {'error': error_msg}

    wb = load_workbook(filename, data_only=True)
    by_type = {err: [] for err in EXCEL_ERRORS}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for err in EXCEL_ERRORS:
                        if err in cell.value:
                            by_type[err].append(f"{ws.title}!{cell.coordinate}")
                            break
    wb.close()

    wb_formulas = load_workbook(filename, data_only=False)
    formula_count = sum(1 for ws in wb_formulas.worksheets for row in ws.iter_rows() for cell in row
                        if isinstance(cell.value, str) and cell.value.startswith('='))
    wb_formulas.close()

    result = summarize(by_type, formula_count, 'legacy-libreoffice-resave')
    result['warning'] = ('paper-xlsx oracle unavailable (stock openpyxl installed); LibreOffice re-saved '
                         'the whole file and may have dropped charts, images, OLE objects or comments')
    return result


def recalc(filename, timeout=30):
    """
    Recalculate formulas in Excel file and report any errors

    Returns:
        dict with error locations and counts
    """
    if not Path(filename).exists():
        return {'error': f'File {filename} does not exist'}
    try:
        from openpyxl import oracle
    except ImportError:
        oracle = None
    try:
        if oracle is None:
            return recalc_legacy(filename, timeout)
        if not oracle.available():
            return {'error': 'LibreOffice not found (looked on PATH and in /Applications/LibreOffice.app)',
                    'method': 'paper-oracle'}
        return recalc_oracle(filename, timeout, oracle)
    except Exception as e:
        return {'error': f'{type(e).__name__}: {e}'}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        print("  - method: 'paper-oracle' or 'legacy-libreoffice-resave'")
        sys.exit(1)

    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = recalc(filename, timeout)
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
